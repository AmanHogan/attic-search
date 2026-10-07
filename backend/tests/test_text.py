import sqlite3
from collections.abc import Callable
from pathlib import Path

import pytest

from attic import config, text
from attic.config import Paths
from attic.ingest import ingest_inbox
from attic.text import run_ocr


def job_status(conn: sqlite3.Connection) -> list[str]:
    return [row[0] for row in conn.execute("SELECT status FROM jobs ORDER BY job_id")]


def test_pdf_text_layer_is_used(
    paths: Paths, conn: sqlite3.Connection, make_pdf: Callable[..., Path]
) -> None:
    """A PDF that already contains text skips OCR and becomes searchable.

    Args:
        paths (Paths): Temporary archive locations.
        conn (sqlite3.Connection): Initialized test database.
        make_pdf (Callable[..., Path]): Test PDF factory.
    """
    make_pdf(paths.inbox / "invoice.pdf", text="Invoice number 12345 from Jiffy Lube")
    ingest_inbox(conn, paths)

    [r] = list(run_ocr(conn, paths))

    assert r.status == "done" and r.engine == "pdf_text" and r.confidence == 1.0
    row = conn.execute("SELECT text_path, is_active FROM page_text").fetchone()
    assert row["is_active"] == 1
    assert "Jiffy Lube" in (paths.archive / row["text_path"]).read_text()
    hits = conn.execute("SELECT page_id FROM pages_fts WHERE pages_fts MATCH 'jiffy'").fetchall()
    assert len(hits) == 1
    assert job_status(conn) == ["done"]


def test_image_uses_apple_vision(
    paths: Paths, conn: sqlite3.Connection, make_image: Callable[..., Path]
) -> None:
    """A photo has no text layer, so Apple Vision reads it.

    Args:
        paths (Paths): Temporary archive locations.
        conn (sqlite3.Connection): Initialized test database.
        make_image (Callable[..., Path]): Test image factory.
    """
    make_image(paths.inbox / "note.png", size=(900, 200), text="HELLO WORLD", font_size=72)
    ingest_inbox(conn, paths)

    [r] = list(run_ocr(conn, paths))

    assert r.status == "done" and r.engine == "vision"
    assert r.confidence is not None and r.confidence > 0
    path = conn.execute("SELECT text_path FROM page_text").fetchone()[0]
    assert "HELLO WORLD" in (paths.archive / path).read_text().upper()


def test_rerun_is_noop(paths: Paths, conn: sqlite3.Connection, make_pdf: Callable[..., Path]) -> None:
    make_pdf(paths.inbox / "a.pdf", text="some text that is long enough")
    ingest_inbox(conn, paths)
    list(run_ocr(conn, paths))

    assert list(run_ocr(conn, paths)) == []
    assert conn.execute("SELECT count(*) FROM page_text").fetchone()[0] == 1


def test_failure_is_recorded_and_run_continues(
    paths: Paths,
    conn: sqlite3.Connection,
    make_image: Callable[..., Path],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """If OCR fails on a page, the job is marked error and nothing is stored for it.

    Args:
        paths (Paths): Temporary archive locations.
        conn (sqlite3.Connection): Initialized test database.
        make_image (Callable[..., Path]): Test image factory.
        monkeypatch (pytest.MonkeyPatch): Replaces Vision OCR with a failing stub.
    """
    make_image(paths.inbox / "a.png", text="one")
    make_image(paths.inbox / "b.png", text="two")
    ingest_inbox(conn, paths)

    def boom(image: Path) -> text.PageText:
        raise RuntimeError("vision exploded")

    monkeypatch.setattr(text, "_vision_text", boom)
    results = list(run_ocr(conn, paths))

    assert [r.status for r in results] == ["error", "error"]
    assert job_status(conn) == ["error", "error"]
    assert conn.execute("SELECT error FROM jobs").fetchone()[0] == "vision exploded"
    assert conn.execute("SELECT count(*) FROM page_text").fetchone()[0] == 0


def test_docx_text_includes_header_and_footer(tmp_path: Path) -> None:
    import docx

    from attic.text import _docx_text

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Prof. Van Vleet"
    document.sections[0].footer.paragraphs[0].text = "Page footer"
    document.add_paragraph("Body line")
    path = tmp_path / "paper.docx"
    document.save(str(path))

    assert _docx_text(path).splitlines() == ["Prof. Van Vleet", "Page footer", "Body line"]


def test_requeue_word_rereads_only_word_files(
    paths: Paths, conn: sqlite3.Connection, make_pdf: Callable[..., Path]
) -> None:
    import docx

    from attic.ingest import ingest_file
    from attic.text import requeue_word

    word = docx.Document()
    word.add_paragraph("body")
    word_path = paths.inbox / "a.docx"
    word.save(str(word_path))
    ingest_file(conn, paths, word_path)
    ingest_file(conn, paths, make_pdf(paths.inbox / "b.pdf", text="some long text " * 4))
    list(run_ocr(conn, paths))

    assert requeue_word(conn) == 1
    assert sorted(job_status(conn)) == ["done", "pending"]


def test_sparse_page_without_image_keeps_its_text(paths: Paths, conn: sqlite3.Connection) -> None:
    """A near-blank page of a long book (no image to OCR) is stored, not an error."""
    import pymupdf

    pdf = paths.inbox / "book.pdf"
    doc = pymupdf.open()
    for i in range(config.LONG_PDF_MIN_PAGES):
        page = doc.new_page()
        if i != 10:  # one blank page, past the rendered ones
            page.insert_text((72, 72), "chapter text goes on and on here " * 3, fontsize=11)
    doc.save(pdf)
    doc.close()
    ingest_inbox(conn, paths)

    results = list(run_ocr(conn, paths))

    assert all(r.status == "done" for r in results)


def _make_deck(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Intro to Compilers"
    first.placeholders[1].text = "Lexing and parsing"
    first.notes_slide.notes_text_frame.text = "Mention the dragon book"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Grammar table"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for r, row in enumerate((("symbol", "rule"), ("expr", "term + term"))):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    deck.save(str(path))
    return path


def test_pptx_is_one_text_page_per_slide(paths: Paths, conn: sqlite3.Connection) -> None:
    """A deck becomes one imageless page per slide, with notes and table cells in the text."""
    _make_deck(paths.inbox / "lecture.pptx")

    [result] = ingest_inbox(conn, paths)
    assert (result.status, result.pages) == ("ingested", 2)
    assert conn.execute("SELECT count(*) FROM pages WHERE image_path IS NOT NULL").fetchone()[0] == 0

    assert all(r.status == "done" and r.engine == "pptx" for r in run_ocr(conn, paths))
    texts = [r[0] for r in conn.execute("SELECT text FROM pages_fts ORDER BY rowid")]
    assert "Intro to Compilers" in texts[0] and "Mention the dragon book" in texts[0]
    assert "term + term" in texts[1]


def test_broken_pptx_is_an_error(paths: Paths, conn: sqlite3.Connection) -> None:
    (paths.inbox / "bad.pptx").write_bytes(b"not a deck")

    [result] = ingest_inbox(conn, paths)

    assert result.status == "error"
