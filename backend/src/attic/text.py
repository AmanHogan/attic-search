"""Extract text for each page: the PDF's own text layer when it has one, else Apple Vision OCR.

Works through pending `ocr` jobs queued by ingest. Every result is kept in
`page_text` (never overwritten by a different engine/version), the newest is
marked active, and the active text is what goes into the FTS5 keyword index.
"""

import sqlite3
import subprocess
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
import pymupdf
from ocrmac import ocrmac
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from attic import config
from attic.code import notebook_text
from attic.config import Paths
from attic.db import now


@dataclass(frozen=True)
class PageText:
    """Text extracted from one page by one engine.

    Attributes:
        text (str): The page's text, one line per line found.
        engine (str): "pdf_text" (a PDF's own text layer), "docx" (`python-docx`), "pptx"
            (`python-pptx`, one slide), "textutil" (legacy .doc via macOS), "vision" (Apple
            Vision OCR), or "source" (a code file).
        confidence (float): 0-1. Always 1.0 for a PDF text layer.
    """

    text: str
    engine: str
    confidence: float


@dataclass(frozen=True)
class OcrResult:
    """Outcome of one OCR job.

    Attributes:
        source_filename (str): Original name of the raw file the page came from.
        page_no (int): 1-based page number within that raw file.
        status (str): "done" or "error".
        engine (str | None): Engine that produced the text; None on error.
        confidence (float | None): Page confidence; None on error.
        chars (int): Length of the extracted text.
        error (str | None): Why extraction failed; set only on error.
    """

    source_filename: str
    page_no: int
    status: str
    engine: str | None = None
    confidence: float | None = None
    chars: int = 0
    error: str | None = None


# --- start private functions ---


def _pdf_text(pdf: Path, page_no: int) -> str:
    """Read the text layer of one PDF page.

    Args:
        pdf (Path): PDF file.
        page_no (int): 1-based page number.

    Returns:
        str: The page's embedded text, stripped; empty if it has none.
    """
    with pymupdf.open(pdf) as doc:
        return str(doc[page_no - 1].get_text()).strip()


def _source_text(source: Path) -> str:
    """Read a source file as text; a notebook is reduced to its code and markdown cells.

    Args:
        source (Path): Source file.

    Returns:
        str: The file's text, cut to `CODE_MAX_BYTES` characters (a notebook can be large
            even after its outputs are dropped).
    """
    is_notebook = source.suffix.lower() == ".ipynb"
    text = notebook_text(source) if is_notebook else source.read_text(errors="replace")
    return text.strip()[: config.CODE_MAX_BYTES]


def _shape_lines(shapes: Iterable[Any]) -> list[str]:
    """Collect the text from slide shapes: text frames, table rows, and grouped shapes.

    Args:
        shapes (Iterable[Any]): A slide's (or group's) shapes.

    Returns:
        list[str]: One entry per paragraph or table row, in reading order of the shapes.
    """
    lines: list[str] = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_lines(shape.shapes))
        elif shape.has_text_frame:
            lines.extend(p.text.strip() for p in shape.text_frame.paragraphs)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                lines.append("\t".join(cell.text.strip() for cell in row.cells))
    return lines


def _pptx_text(deck: Path, slide_no: int) -> str:
    """Read one slide's text, then its speaker notes, with `python-pptx`.

    Text inside pictures is not read.

    Args:
        deck (Path): .pptx file.
        slide_no (int): 1-based slide number.

    Returns:
        str: The slide's text, one paragraph or table row per line.
    """
    slide = Presentation(str(deck)).slides[slide_no - 1]
    lines = _shape_lines(slide.shapes)
    if slide.has_notes_slide:
        lines.append(slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(line for line in lines if line).strip()


def _docx_text(doc: Path) -> str:
    """Read a .docx file's text with `python-docx`.

    Page headers and footers come first, since a student's name or a running head
    lives there and `python-docx` keeps it out of `paragraphs`. Table cells are
    included after the paragraphs for the same reason: a lab report's data would
    otherwise be lost. Header text repeated across sections is kept once.

    Args:
        doc (Path): .docx file.

    Returns:
        str: The document's text, one paragraph or table row per line.
    """
    document = docx.Document(str(doc))
    lines: list[str] = []
    for section in document.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            lines.extend(p.text.strip() for p in part.paragraphs if p.text.strip() not in lines)
    lines.extend(p.text.strip() for p in document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(line for line in lines if line).strip()


def _doc_text(doc: Path) -> str:
    """Read a legacy .doc file's text with macOS `textutil`.

    `python-docx` cannot open the pre-2007 binary format, and no portable library
    reads it well, so this one path is macOS-only.

    Args:
        doc (Path): .doc file.

    Returns:
        str: The document's text.

    Raises:
        RuntimeError: If `textutil` is missing or cannot read the file.
    """
    if not Path(config.TEXTUTIL).exists():
        raise RuntimeError(f"reading {config.DOC_SUFFIX} needs macOS {config.TEXTUTIL}; re-save as .docx")
    result = subprocess.run(
        [config.TEXTUTIL, "-convert", "txt", "-stdout", str(doc)],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(f"textutil failed: {result.stderr.strip() or result.returncode}")
    return result.stdout.strip()


def _vision_text(image: Path) -> PageText:
    """Run Apple Vision OCR on a page image.

    The page confidence is the average of the per-line confidences,
    weighted by line length so a long, well-read line counts for more
    than a stray two-letter fragment.

    Args:
        image (Path): Page image.

    Returns:
        PageText: Recognized lines joined with newlines, engine "vision".
    """
    lines = ocrmac.OCR(
        str(image), recognition_level="accurate", language_preference=list(config.OCR_LANGUAGES)
    ).recognize()
    chars = sum(len(text) for text, _, _ in lines)
    confidence = sum(len(text) * conf for text, conf, _ in lines) / chars if chars else 0.0
    return PageText("\n".join(text for text, _, _ in lines), "vision", float(confidence))


def _store_text(conn: sqlite3.Connection, paths: Paths, job: sqlite3.Row, result: PageText) -> None:
    """Save extracted text to disk, `page_text`, and FTS, and mark the job done.

    The document's index job is set back to pending, since its chunks were built
    from the text this replaces.

    All database writes happen in one transaction, so a crash leaves the job
    pending rather than half-recorded.

    Args:
        conn (sqlite3.Connection): Open archive database.
        paths (Paths): Archive locations.
        job (sqlite3.Row): Pending job joined with its page (see `run_ocr`).
        result (PageText): Text to store.
    """
    version = config.OCR_STAGE_VERSION
    text_path = f"derived/{job['raw_sha256']}/p{job['raw_page_no']:03d}.{result.engine}.v{version}.md"
    out = paths.archive / text_path
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(result.text + "\n")

    ts = now()
    with conn:
        conn.execute("UPDATE page_text SET is_active = 0 WHERE page_id = ?", (job["page_id"],))
        conn.execute(
            "INSERT OR REPLACE INTO page_text"
            " (page_id, engine, version, confidence, text_path, is_active, created_at)"
            " VALUES (?, ?, ?, ?, ?, 1, ?)",
            (job["page_id"], result.engine, version, result.confidence, text_path, ts),
        )
        conn.execute("DELETE FROM pages_fts WHERE page_id = ?", (job["page_id"],))
        conn.execute(
            "INSERT INTO pages_fts (text, page_id, doc_id) VALUES (?, ?, ?)",
            (result.text, job["page_id"], job["doc_id"]),
        )
        # Chunks are built from this text, so new text means the document needs re-indexing.
        conn.execute(
            "UPDATE jobs SET status = 'pending', updated_at = ? WHERE target_id = ? AND stage = ?",
            (ts, job["doc_id"], config.INDEX_STAGE),
        )
        conn.execute(
            "UPDATE jobs SET status = 'done', error = NULL, updated_at = ? WHERE job_id = ?",
            (ts, job["job_id"]),
        )


# --- end private functions ---


def extract_text(paths: Paths, raw_path: str, raw_page_no: int, image_path: str | None) -> PageText:
    """Get a page's text, preferring a file's own text over OCR.

    Args:
        paths (Paths): Archive locations.
        raw_path (str): Raw file, relative to `archive/`.
        raw_page_no (int): 1-based page number within the raw file.
        image_path (str | None): Rendered page image, relative to `archive/`;
            None for a Word file, which has no image.

    Returns:
        PageText: The extracted text and which engine produced it.
    """
    suffix = Path(raw_path).suffix.lower()
    if suffix in config.CODE_SUFFIXES:
        return PageText(_source_text(paths.archive / raw_path), "source", 1.0)
    if suffix == config.PPTX_SUFFIX:
        return PageText(_pptx_text(paths.archive / raw_path, raw_page_no), "pptx", 1.0)
    if suffix == config.DOCX_SUFFIX:
        return PageText(_docx_text(paths.archive / raw_path), "docx", 1.0)
    if suffix == config.DOC_SUFFIX:
        return PageText(_doc_text(paths.archive / raw_path), "textutil", 1.0)
    if raw_path.endswith(config.PDF_SUFFIX):
        layer = _pdf_text(paths.archive / raw_path, raw_page_no)
        # A page of a long book has no image, so a sparse text layer (a blank or chapter-title
        # page) is all there is; keep it rather than fail and block the whole book.
        if image_path is None or len("".join(layer.split())) >= config.TEXT_LAYER_MIN_CHARS:
            return PageText(layer, "pdf_text", 1.0)
    if image_path is None:
        raise RuntimeError(f"no page image to read text from: {raw_path}")
    return _vision_text(paths.archive / image_path)


def requeue_word(conn: sqlite3.Connection) -> int:
    """Mark every Word file's OCR job pending, so its text is read again.

    Use after improving how Word text is extracted; `run_ocr` then replaces the
    stored text, keyword index, and (via `_store_text`) re-queues indexing.

    Args:
        conn (sqlite3.Connection): Open archive database.

    Returns:
        int: Number of pages queued.
    """
    marks = ", ".join("?" * len(config.WORD_SUFFIXES))
    with conn:
        cur = conn.execute(
            "UPDATE jobs SET status = 'pending', error = NULL, updated_at = ?"
            " WHERE stage = ? AND version = ? AND target_id IN ("
            "   SELECT p.page_id FROM pages p JOIN raw_files r ON r.sha256 = p.raw_sha256"
            f"   WHERE lower(substr(r.raw_path, instr(r.raw_path, '.'))) IN ({marks}))",
            (now(), config.OCR_STAGE, config.OCR_STAGE_VERSION, *sorted(config.WORD_SUFFIXES)),
        )
    return cur.rowcount


def run_ocr(conn: sqlite3.Connection, paths: Paths) -> Iterator[OcrResult]:
    """Process every pending OCR job, yielding each result as soon as it's done.

    Failures are recorded on the job (status "error") and don't stop the run.

    Args:
        conn (sqlite3.Connection): Open archive database.
        paths (Paths): Archive locations.

    Yields:
        OcrResult: Outcome for each page, in the order the jobs were queued.
    """
    jobs = conn.execute(
        "SELECT j.job_id, p.page_id, p.doc_id, p.raw_sha256, p.raw_page_no, p.image_path,"
        " r.raw_path, r.source_filename"
        " FROM jobs j"
        " JOIN pages p ON p.page_id = j.target_id"
        " JOIN raw_files r ON r.sha256 = p.raw_sha256"
        " WHERE j.stage = ? AND j.version = ? AND j.status = 'pending'"
        " ORDER BY j.job_id",
        (config.OCR_STAGE, config.OCR_STAGE_VERSION),
    ).fetchall()

    for job in jobs:
        name, page_no = job["source_filename"], job["raw_page_no"]
        try:
            result = extract_text(paths, job["raw_path"], page_no, job["image_path"])
            _store_text(conn, paths, job, result)
        except Exception as e:
            with conn:
                conn.execute(
                    "UPDATE jobs SET status = 'error', error = ?, updated_at = ? WHERE job_id = ?",
                    (str(e), now(), job["job_id"]),
                )
            yield OcrResult(name, page_no, "error", error=str(e))
            continue
        yield OcrResult(name, page_no, "done", result.engine, result.confidence, len(result.text))
